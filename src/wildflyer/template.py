"""Generate input workbooks: the blank template, or one pre-filled with data."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .catalog import CATALOG, ROLES
from .output import no_formulas
from .loader import (LOCK_COLUMNS, RULE_COLUMNS, SETTINGS_KEYS, SHEET_LOCKS, SHEET_RULES, SHEET_SETTINGS,
                     SHEET_TEAMS, TEAM_COLUMNS)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical="top")

RULE_WIDTHS = {"ID": 8, "Enabled": 9, "Type": 26, "Teams": 16, "Opponents": 14, "Role": 8,
               "From": 10, "To": 10, "Dates": 30, "Days": 12, "Min": 6, "Max": 6, "N": 6, "Option": 10,
               "Hard/Soft": 10, "Weight": 8, "Note": 50}
SETTINGS_HELP = {
    "Run name": "Short label for this run (used in the output folder name)",
    "Season start": "First date games may be scheduled (a full date, e.g. 2027-06-12)",
    "Season end": "Last date games may be scheduled",
    "Time limit (seconds)": "How long the solver may search (default 300)",
    "Base run": "Optional: a previous schedule workbook (from wildflyer solve) to build from and compare "
                "against, e.g. schedule_v3.xlsx (relative to this file's folder)",
    "Change weight": "Optional: soft cost for each base-run game that is moved or dropped (blank/0 = off)",
    "Lock base before": "Optional: freeze the schedule before this date exactly as in the base run "
                        "(same games, none added)",
}


def build_workbook(
    settings: dict[str, Any] | None = None,
    teams: Iterable[Sequence[Any]] = (),
    rules: Iterable[dict[str, Any]] = (),
    locks: Iterable[Sequence[Any]] = (),
) -> Workbook:
    """Build an input workbook. `rules` are dicts keyed by Rules column name."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_SETTINGS
    _settings_sheet(ws, settings or {})
    _table_sheet(wb.create_sheet(SHEET_TEAMS), TEAM_COLUMNS, teams, {"Code": 8, "Name": 28, "Color": 10})
    _color_team_rows(wb[SHEET_TEAMS])
    rules_ws = wb.create_sheet(SHEET_RULES)
    _table_sheet(rules_ws, RULE_COLUMNS, ([r.get(c) for c in RULE_COLUMNS] for r in rules), RULE_WIDTHS)
    _text_column(rules_ws, RULE_COLUMNS.index("Dates") + 1)
    _table_sheet(wb.create_sheet(SHEET_LOCKS), LOCK_COLUMNS, locks, {"Date": 12, "Home": 8, "Away": 8})
    _help_sheet(wb.create_sheet("Help"))
    _rule_validations(wb, rules_ws)
    no_formulas(wb)
    return wb


def write_template(path: str | Path, **data: Any) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(**data).save(path)
    return path


# --- sheets -----------------------------------------------------------------------


def _header(ws: Worksheet, columns: Sequence[str], widths: dict[str, int]) -> None:
    ws.append(list(columns))
    for i, name in enumerate(columns, start=1):
        cell = ws.cell(1, i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 12)
    ws.freeze_panes = "A2"


def _table_sheet(ws: Worksheet, columns: Sequence[str], rows: Iterable[Sequence[Any]],
                 widths: dict[str, int]) -> None:
    _header(ws, columns, widths)
    for row in rows:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"


def _text_column(ws: Worksheet, col: int, last_row: int = 1000) -> None:
    """Format a column as Text so Excel doesn't turn entries like 6/25 into dates."""
    for r in range(2, last_row + 1):
        ws.cell(r, col).number_format = "@"


def _color_team_rows(ws: Worksheet) -> None:
    col = TEAM_COLUMNS.index("Color") + 1
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell = row[0]
        if isinstance(cell.value, str) and len(cell.value.lstrip("#")) == 6:
            cell.fill = PatternFill("solid", fgColor=cell.value.lstrip("#").upper())


def _settings_sheet(ws: Worksheet, settings: dict[str, Any]) -> None:
    _header(ws, ("Setting", "Value", "Help"), {"Setting": 22, "Value": 22, "Help": 80})
    for label in SETTINGS_KEYS.values():
        ws.append([label, settings.get(label), SETTINGS_HELP[label]])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(row[0].value, date):
            row[0].number_format = "yyyy-mm-dd"


def _help_sheet(ws: Worksheet) -> None:
    ws.column_dimensions["A"].width = 28
    for col, width in zip("BCDEFG", (22, 70, 46, 18, 42, 26)):
        ws.column_dimensions[col].width = width

    def section(title: str) -> None:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = SECTION_FONT

    ws.append(["Wildflyer input workbook"])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append(["Fill in Settings, Teams and Rules (Locks is optional), then run "
               "`wildflyer validate <file>` to check it."])

    section("Writing rules")
    for line in (
        ("ID", "Unique label, used in reports (e.g. R01)."),
        ("Enabled", "Y or N (blank = Y). N keeps the rule but skips it for this run."),
        ("Hard/Soft", "Hard = must hold. Soft = may be broken at a cost of Weight per violation."),
        ("Weight", "Soft rules only: cost per violation unit (see the table below). Higher = more important."),
        ("Teams / Opponents", "Team codes: CHI  or  CHI, UTA  or  ALL  or  ALL except CHI, UTA"),
        ("From / To", "Locations, named by the hosting team's code (every game is at the home team's "
                      "location). Also: any, home (the team's own home games), away (games hosted by "
                      "anyone else)."),
        ("Role", "home, away or any (blank = any)."),
        ("Dates", "6/25  ·  6/25-6/27  ·  6/25-27  ·  2027-06-25  ·  8/2-end  ·  start-6/20  ·  "
                  "lists: 6/25-27, 7/1-3. Blank = whole season."),
        ("Days", "Mon-Thu  ·  Fri,Sat,Sun  ·  Wed. Blank = every day."),
        ("Min / Max / N / Option", "Meaning depends on the rule type (below)."),
    ):
        ws.append(list(line))
        ws.cell(ws.max_row, 2).alignment = WRAP

    section("Definitions")
    for line in (
        ("Off day", "A date on which the team has no game."),
        ("Block", "A run of a team's games against one opponent with no game against anyone else in "
                  "between. Off days do not break a block."),
        ("Move", "Two consecutive games of a team at different locations (different host teams)."),
        ("Homestand / road trip", "A run of consecutive home (or away) games of a team."),
    ):
        ws.append(list(line))
        ws.cell(ws.max_row, 2).alignment = WRAP

    section("Rule types")
    header = ["Type", "Category", "What it does", "Fields (* = required)", "Option", "Example",
              "Soft violation unit"]
    ws.append(header)
    for i in range(1, len(header) + 1):
        ws.cell(ws.max_row, i).fill, ws.cell(ws.max_row, i).font = HEADER_FILL, HEADER_FONT
    for spec in CATALOG.values():
        fields = ", ".join(f"{_title(f)}{'*' if fs.required else ''}" for f, fs in spec.fields.items())
        option = " / ".join(spec.options)
        if spec.options:
            option += f" (default {spec.option_default})" if spec.option_default else " (required)"
        summary = spec.summary + (" Soft only." if spec.soft_only else "")
        ws.append([spec.name, spec.category, summary, fields, option, spec.example, spec.violation_unit])
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP


def _title(field_name: str) -> str:
    return {"n": "N", "from": "From", "to": "To"}.get(field_name, field_name.capitalize())


def _rule_validations(wb: Workbook, ws: Worksheet) -> None:
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    options = sorted({o for s in CATALOG.values() for o in s.options})
    columns = {"Type": list(CATALOG), "Enabled": ["Y", "N"], "Hard/Soft": ["Hard", "Soft"],
               "Role": list(ROLES), "Option": options}
    for col_idx, (name, values) in enumerate(columns.items(), start=1):
        letter = get_column_letter(col_idx)
        lists.cell(1, col_idx, name)
        for i, v in enumerate(values, start=2):
            lists.cell(i, col_idx, v)
        target = get_column_letter(RULE_COLUMNS.index(name) + 1)
        dv = DataValidation(type="list", formula1=f"Lists!${letter}$2:${letter}${len(values) + 1}",
                            allow_blank=True, showErrorMessage=name != "Option")
        dv.add(f"{target}2:{target}1000")
        ws.add_data_validation(dv)
