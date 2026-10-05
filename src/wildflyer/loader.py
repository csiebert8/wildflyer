"""Read an input workbook into a validated RunInput, collecting row-level issues."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from . import selectors as sel
from .catalog import CATALOG, FIELDS, RuleSpec
from .model import Issue, Level, Lock, Rule, RunInput, Settings, Team, Venue

SHEET_SETTINGS = "Settings"
SHEET_TEAMS = "Teams"
SHEET_VENUES = "Venues"
SHEET_RULES = "Rules"
SHEET_LOCKS = "Locks"

TEAM_COLUMNS = ("Code", "Name", "Home venue", "Color")
VENUE_COLUMNS = ("Code", "Name", "Team")
RULE_COLUMNS = ("ID", "Enabled", "Type", "Teams", "Opponents", "Role", "Venues", "From", "To",
                "Dates", "Days", "Min", "Max", "N", "Option", "Hard/Soft", "Weight", "Note")
LOCK_COLUMNS = ("Date", "Home", "Away", "Venue")
SETTINGS_KEYS = {
    "runname": "Run name",
    "seasonstart": "Season start",
    "seasonend": "Season end",
    "timelimitseconds": "Time limit (seconds)",
    "baserun": "Base run",
    "changeweight": "Change weight",
}

_YES = {"y", "yes", "true", "1", "on"}
_NO = {"n", "no", "false", "0", "off"}


@dataclass
class LoadResult:
    run: RunInput | None
    issues: list[Issue] = field(default_factory=list)

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level is Level.ERROR]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.level is Level.WARNING]

    @property
    def ok(self) -> bool:
        return not self.errors


def load(path: str | Path) -> LoadResult:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # unreadable / not an xlsx
        return LoadResult(None, [Issue(Level.ERROR, str(path), f"can't open workbook: {e}")])
    return _Loader(wb).run()


def _norm(text: Any) -> str:
    return re.sub(r"[^a-z0-9/]", "", str(text or "").lower())


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


class _Loader:
    def __init__(self, wb: openpyxl.Workbook):
        self.wb = wb
        self.issues: list[Issue] = []

    def _issue(self, level: Level, sheet: str, msg: str, **kw: Any) -> None:
        self.issues.append(Issue(level, sheet, msg, **kw))

    def run(self) -> LoadResult:
        settings = self._settings()
        venues = self._venues()
        teams = self._teams(venues)
        rules: list[Rule] = []
        locks: list[Lock] = []
        if settings is None:
            self._issue(Level.ERROR, SHEET_RULES, "rules not checked until Settings season dates are fixed")
        else:
            rules = self._rules(settings, teams, venues)
            locks = self._locks(settings, teams, venues)
        if any(i.level is Level.ERROR for i in self.issues) or settings is None:
            return LoadResult(None, self.issues)
        return LoadResult(RunInput(settings, teams, venues, rules, locks), self.issues)

    # --- sheet helpers ---------------------------------------------------------

    def _table(self, name: str, columns: tuple[str, ...], required: bool = True):
        """Yield (row_number, {column: value}) for non-blank rows of a header table."""
        if name not in self.wb.sheetnames:
            if required:
                self._issue(Level.ERROR, name, "sheet is missing")
            return []
        ws = self.wb[name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            if required:
                self._issue(Level.ERROR, name, "sheet is empty (expected a header row)")
            return []
        index = {_norm(h): i for i, h in enumerate(header) if not _blank(h)}
        missing = [c for c in columns if _norm(c) not in index]
        if missing:
            self._issue(Level.ERROR, name, f"missing column(s): {', '.join(missing)}", row=1)
            return []
        known = {_norm(c) for c in columns}
        extra = [str(h) for h in header if not _blank(h) and _norm(h) not in known]
        if extra:
            self._issue(Level.WARNING, name, f"ignoring unknown column(s): {', '.join(extra)}", row=1)
        out = []
        for r, values in enumerate(rows, start=2):
            record = {c: (values[index[_norm(c)]] if index[_norm(c)] < len(values) else None) for c in columns}
            if all(_blank(v) for v in record.values()):
                continue
            out.append((r, record))
        return out

    # --- Settings ------------------------------------------------------------------

    def _settings(self) -> Settings | None:
        if SHEET_SETTINGS not in self.wb.sheetnames:
            self._issue(Level.ERROR, SHEET_SETTINGS, "sheet is missing")
            return None
        values: dict[str, tuple[int, Any]] = {}
        for r, row in enumerate(self.wb[SHEET_SETTINGS].iter_rows(values_only=True), start=1):
            if not row or _blank(row[0]):
                continue
            key = _norm(row[0]).replace("/", "")
            if key in ("setting", "settings"):  # header row
                continue
            if key not in SETTINGS_KEYS:
                self._issue(Level.WARNING, SHEET_SETTINGS, f"unknown setting '{row[0]}' ignored", row=r)
                continue
            values[key] = (r, row[1] if len(row) > 1 else None)

        def get(key: str, parse, required: bool = False, default: Any = None) -> Any:
            r, v = values.get(key, (None, None))
            if _blank(v):
                if required:
                    self._issue(Level.ERROR, SHEET_SETTINGS, f"'{SETTINGS_KEYS[key]}' is required",
                                row=r, field=SETTINGS_KEYS[key])
                return default
            try:
                return parse(v)
            except ValueError as e:
                self._issue(Level.ERROR, SHEET_SETTINGS, str(e), row=r, field=SETTINGS_KEYS[key])
                return default

        start = get("seasonstart", sel.parse_date_cell, required=True)
        end = get("seasonend", sel.parse_date_cell, required=True)
        time_limit = get("timelimitseconds", sel.parse_int, default=300)
        change_weight = get("changeweight", sel.parse_number, default=0.0)
        run_name = get("runname", _text, default="")
        base_run = get("baserun", _text, default=None)
        if time_limit is not None and time_limit <= 0:
            self._issue(Level.ERROR, SHEET_SETTINGS, "time limit must be positive", field="Time limit (seconds)")
        if change_weight is not None and change_weight < 0:
            self._issue(Level.ERROR, SHEET_SETTINGS, "change weight can't be negative", field="Change weight")
        if start is None or end is None:
            return None
        if end < start:
            self._issue(Level.ERROR, SHEET_SETTINGS, "Season end is before Season start", field="Season end")
            return None
        if (end - start).days > 400:
            self._issue(Level.WARNING, SHEET_SETTINGS, "season is longer than 400 days; check the dates")
        return Settings(start, end, run_name, time_limit or 300, base_run, change_weight or 0.0)

    # --- Teams & venues -------------------------------------------------------------

    def _venues(self) -> dict[str, Venue]:
        venues: dict[str, Venue] = {}
        self._venue_owner_rows: dict[str, tuple[int, str]] = {}
        for r, rec in self._table(SHEET_VENUES, VENUE_COLUMNS):
            code = _text(rec["Code"]).upper()
            if not code:
                self._issue(Level.ERROR, SHEET_VENUES, "Code is required", row=r, field="Code")
                continue
            if not re.fullmatch(r"[A-Z0-9_]+", code):
                self._issue(Level.ERROR, SHEET_VENUES, f"code '{code}' may only use letters, digits and _",
                            row=r, field="Code")
                continue
            if code in venues:
                self._issue(Level.ERROR, SHEET_VENUES, f"duplicate venue code '{code}'", row=r, field="Code")
                continue
            owner = _text(rec["Team"]).upper() or None
            if owner:
                self._venue_owner_rows[code] = (r, owner)
            venues[code] = Venue(code, _text(rec["Name"]) or code, owner)
        if not venues and SHEET_VENUES in self.wb.sheetnames:
            self._issue(Level.ERROR, SHEET_VENUES, "no venues defined")
        return venues

    def _teams(self, venues: dict[str, Venue]) -> dict[str, Team]:
        teams: dict[str, Team] = {}
        for r, rec in self._table(SHEET_TEAMS, TEAM_COLUMNS):
            code = _text(rec["Code"]).upper()
            if not code:
                self._issue(Level.ERROR, SHEET_TEAMS, "Code is required", row=r, field="Code")
                continue
            if not re.fullmatch(r"[A-Z0-9_]+", code) or code == "ALL":
                self._issue(Level.ERROR, SHEET_TEAMS, f"code '{code}' must use letters, digits and _ "
                            "and can't be ALL", row=r, field="Code")
                continue
            if code in teams:
                self._issue(Level.ERROR, SHEET_TEAMS, f"duplicate team code '{code}'", row=r, field="Code")
                continue
            home = _text(rec["Home venue"]).upper()
            if not home:
                self._issue(Level.ERROR, SHEET_TEAMS, "Home venue is required", row=r, field="Home venue")
            elif venues and home not in venues:
                self._issue(Level.ERROR, SHEET_TEAMS, f"home venue '{home}' is not on the Venues sheet",
                            row=r, field="Home venue")
            color = _text(rec["Color"]).lstrip("#").upper()
            if not color:
                color = "D9D9D9"
            elif not re.fullmatch(r"[0-9A-F]{6}", color):
                self._issue(Level.ERROR, SHEET_TEAMS, f"color '{rec['Color']}' must be a hex code like 76D6FF",
                            row=r, field="Color")
            teams[code] = Team(code, _text(rec["Name"]) or code, home, color)
        if len(teams) < 2 and SHEET_TEAMS in self.wb.sheetnames:
            self._issue(Level.ERROR, SHEET_TEAMS, "at least two teams are required")

        for vcode, (r, owner) in self._venue_owner_rows.items():
            if owner not in teams:
                self._issue(Level.ERROR, SHEET_VENUES, f"venue owner '{owner}' is not a team", row=r, field="Team")
            elif teams[owner].home_venue != vcode:
                self._issue(Level.WARNING, SHEET_VENUES, f"venue owned by {owner}, but {owner}'s home venue "
                            f"is {teams[owner].home_venue}", row=r, field="Team")
        homes: dict[str, str] = {}
        for t in teams.values():
            if t.home_venue in homes:
                self._issue(Level.WARNING, SHEET_TEAMS, f"{t.code} and {homes[t.home_venue]} share home venue "
                            f"{t.home_venue}; only one of them can host on any date")
            homes.setdefault(t.home_venue, t.code)
        return teams

    # --- Rules -------------------------------------------------------------------

    def _rules(self, settings: Settings, teams: dict[str, Team], venues: dict[str, Venue]) -> list[Rule]:
        rules: list[Rule] = []
        seen_ids: dict[str, int] = {}
        for r, rec in self._table(SHEET_RULES, RULE_COLUMNS):
            rule, enabled = _RuleRow(self, r, rec, settings, teams, venues, seen_ids).parse()
            if rule is not None and enabled:
                rules.append(rule)
        if not rules and SHEET_RULES in self.wb.sheetnames:
            self._issue(Level.WARNING, SHEET_RULES, "no enabled rules")
        return rules

    # --- Locks --------------------------------------------------------------------

    def _locks(self, settings: Settings, teams: dict[str, Team], venues: dict[str, Venue]) -> list[Lock]:
        locks: list[Lock] = []
        for r, rec in self._table(SHEET_LOCKS, LOCK_COLUMNS, required=False):
            errors_before = len(self.issues)
            try:
                d = sel.parse_date_cell(rec["Date"])
                if not settings.season_start <= d <= settings.season_end:
                    self._issue(Level.ERROR, SHEET_LOCKS, f"{d} is outside the season", row=r, field="Date")
            except ValueError as e:
                self._issue(Level.ERROR, SHEET_LOCKS, str(e), row=r, field="Date")
                d = None
            home, away = _text(rec["Home"]).upper(), _text(rec["Away"]).upper()
            for col, code in (("Home", home), ("Away", away)):
                if code not in teams:
                    self._issue(Level.ERROR, SHEET_LOCKS, f"unknown team '{code}'", row=r, field=col)
            if home and home == away:
                self._issue(Level.ERROR, SHEET_LOCKS, "a team can't play itself", row=r)
            venue = _text(rec["Venue"]).upper() or None
            if venue and venue not in venues:
                self._issue(Level.ERROR, SHEET_LOCKS, f"unknown venue '{venue}'", row=r, field="Venue")
            if len(self.issues) == errors_before and d is not None:
                locks.append(Lock(r, d, home, away, venue))
        return locks


class _RuleRow:
    """Parses and validates one Rules row."""

    def __init__(self, loader: _Loader, row: int, rec: dict[str, Any], settings: Settings,
                 teams: dict[str, Team], venues: dict[str, Venue], seen_ids: dict[str, int]):
        self.loader, self.row, self.rec = loader, row, rec
        self.settings, self.teams, self.venues, self.seen_ids = settings, teams, venues, seen_ids
        self.rule_id = _text(rec["ID"])
        self.enabled = True
        self.failed = False

    def issue(self, msg: str, field: str | None = None, warning: bool = False) -> None:
        # Problems on disabled rows are reported as warnings so they don't block a run.
        level = Level.WARNING if warning or not self.enabled else Level.ERROR
        if level is Level.ERROR:
            self.failed = True
        prefix = "" if self.enabled or warning else "(disabled rule) "
        self.loader._issue(level, SHEET_RULES, prefix + msg, row=self.row,
                           rule_id=self.rule_id or None, field=field)

    def parse(self) -> tuple[Rule | None, bool]:
        rec = self.rec
        enabled_text = _text(rec["Enabled"]).lower()
        if enabled_text in _NO:
            self.enabled = False
        elif enabled_text and enabled_text not in _YES:
            self.issue(f"Enabled must be Y or N, got '{rec['Enabled']}'", "Enabled")

        if not self.rule_id:
            self.rule_id = f"row{self.row}"
            self.issue(f"ID is blank; using '{self.rule_id}'", "ID", warning=True)
        if self.rule_id in self.seen_ids:
            self.issue(f"duplicate ID (also on row {self.seen_ids[self.rule_id]})", "ID")
        else:
            self.seen_ids[self.rule_id] = self.row

        type_name = _text(rec["Type"]).upper()
        spec = CATALOG.get(type_name)
        if spec is None:
            self.issue(f"unknown rule type '{rec['Type']}'" if type_name else "Type is required", "Type")
            return None, self.enabled

        values = {f: rec[_column(f)] for f in FIELDS}
        for f, v in values.items():
            if f == "option":  # checked against spec.options in _option()
                continue
            col = _column(f)
            if f not in spec.fields and not _blank(v):
                self.issue(f"{col} is not used by {spec.name}; ignored", col, warning=True)
            elif f in spec.fields and spec.fields[f].required and _blank(v):
                self.issue(f"{col} is required for {spec.name}", col)

        kw: dict[str, Any] = {}
        team_codes = set(self.teams)
        venue_codes = set(self.venues)

        def parse(f: str, fn, *args):
            if f not in spec.fields or _blank(values[f]):
                return None
            try:
                return fn(values[f], *args)
            except ValueError as e:
                self.issue(str(e), _column(f))
                return None

        teams = parse("teams", sel.parse_codes, team_codes, "team")
        if teams is None and "teams" in spec.fields and spec.teams_default_all:
            teams = frozenset(team_codes)
        kw["teams"] = teams or frozenset()
        kw["opponents"] = parse("opponents", sel.parse_codes, team_codes, "team") or frozenset()
        kw["venues"] = parse("venues", sel.parse_codes, venue_codes, "venue") or frozenset()
        kw["from_loc"] = parse("from", sel.parse_location, venue_codes)
        kw["to_loc"] = parse("to", sel.parse_location, venue_codes)
        kw["dates"] = parse("dates", sel.parse_dates, self.settings.season_start, self.settings.season_end) \
            or frozenset()
        kw["days"] = parse("days", sel.parse_days) or frozenset()
        for f in ("min", "max", "n"):
            v = parse(f, sel.parse_int)
            if v is not None and v < 0:
                self.issue(f"{_column(f)} can't be negative", _column(f))
            kw[f] = v

        kw["role"] = self._role(spec, values["role"])
        kw["option"] = self._option(spec, values["option"])
        hard, weight = self._hardness(spec)

        self._check_semantics(spec, kw)
        if self.failed:
            return None, self.enabled
        rule = Rule(id=self.rule_id, type=spec.name, row=self.row, hard=hard, weight=weight,
                    note=_text(rec["Note"]), **kw)
        return rule, self.enabled

    def _role(self, spec: RuleSpec, value: Any) -> str:
        if "role" not in spec.fields:
            return "any"
        role = _text(value).lower()
        if not role:
            return "any" if "any" in spec.roles else ""
        if role not in spec.roles:
            self.issue(f"Role must be one of {' / '.join(spec.roles)}, got '{value}'", "Role")
        return role

    def _option(self, spec: RuleSpec, value: Any) -> str | None:
        option = _text(value).lower()
        if not spec.options:
            if option:
                self.issue(f"Option is not used by {spec.name}; ignored", "Option", warning=True)
            return None
        if not option:
            if spec.option_default is None:
                self.issue(f"Option is required for {spec.name}: {' / '.join(spec.options)}", "Option")
            return spec.option_default
        if option not in spec.options:
            self.issue(f"Option must be one of {' / '.join(spec.options)}, got '{value}'", "Option")
        return option

    def _hardness(self, spec: RuleSpec) -> tuple[bool, float]:
        text = _text(self.rec["Hard/Soft"]).lower()
        if not text:
            hard = not spec.soft_only
        elif text in ("hard", "h"):
            hard = True
        elif text in ("soft", "s"):
            hard = False
        else:
            self.issue(f"Hard/Soft must be Hard or Soft, got '{self.rec['Hard/Soft']}'", "Hard/Soft")
            return True, 0.0
        if hard and spec.soft_only:
            self.issue(f"{spec.name} can only be Soft", "Hard/Soft")
        weight = None
        try:
            weight = sel.parse_number(self.rec["Weight"])
        except ValueError as e:
            self.issue(str(e), "Weight")
        if hard:
            if weight is not None:
                self.issue("Weight is ignored on Hard rules", "Weight", warning=True)
            return True, 0.0
        if weight is None:
            self.issue("Soft rules need a Weight (cost per violation)", "Weight")
            return False, 0.0
        if weight <= 0:
            self.issue("Weight must be greater than 0", "Weight")
        return False, weight

    def _check_semantics(self, spec: RuleSpec, kw: dict[str, Any]) -> None:
        lo, hi, n = kw["min"], kw["max"], kw["n"]
        if spec.needs_min_or_max and lo is None and hi is None:
            self.issue(f"{spec.name} needs Min and/or Max", "Min")
        if lo is not None and hi is not None and lo > hi:
            self.issue(f"Min ({lo}) is greater than Max ({hi})", "Min")
        if spec.single_team and len(kw["teams"]) > 1:
            self.issue("Teams must name exactly one team", "Teams")
        if spec.single_opponent and len(kw["opponents"]) > 1:
            self.issue("Opponents must name exactly one team", "Opponents")
        if spec.single_venue and len(kw["venues"]) > 1:
            self.issue("Venues must name exactly one venue", "Venues")
        if "teams" in spec.fields and spec.fields["teams"].required is False and not kw["teams"] \
                and not _blank(self.rec["Teams"]):
            self.issue("Teams selects no teams", "Teams")

        name = spec.name
        if name == "FIXED_GAME" and kw["teams"] and kw["teams"] == kw["opponents"]:
            self.issue("a team can't play itself", "Opponents")
        if name in ("MATCHUP_GAMES", "OPPONENT_BLOCK", "REMATCH_GAP") and kw["teams"] and kw["opponents"] \
                and kw["teams"] == kw["opponents"] and len(kw["teams"]) == 1:
            self.issue("Teams and Opponents are the same single team", "Opponents")
        if name == "GAMES_IN_WINDOW":
            if kw["option"] == "rolling" and not n:
                self.issue("rolling windows need N (window length in days, at least 1)", "N")
            if kw["option"] == "week" and n is not None:
                self.issue("N is ignored for weekly windows", "N", warning=True)
        if name in ("TRAVEL_REST", "REMATCH_GAP", "OPPONENT_CHANGE_REST") and n == 0:
            self.issue("N = 0 has no effect", "N", warning=True)
        if name in ("MAX_CONSECUTIVE_GAME_DAYS", "HOME_AWAY_RUN") and hi == 0:
            self.issue("Max = 0 forbids all games; use AVAILABILITY or LEAGUE_BLACKOUT instead", "Max")
        if name == "DATE_PREFERENCE" and not kw["dates"] and not kw["days"]:
            self.issue("DATE_PREFERENCE needs Dates and/or Days", "Dates")

        start, end = self.settings.season_start, self.settings.season_end
        outside = sorted(d for d in kw["dates"] if not start <= d <= end)
        if outside:
            shown = ", ".join(d.isoformat() for d in outside[:3]) + (" ..." if len(outside) > 3 else "")
            self.issue(f"{len(outside)} date(s) outside the season are ignored: {shown}", "Dates", warning=True)


def _column(field_name: str) -> str:
    return {"n": "N", "from": "From", "to": "To"}.get(field_name, field_name.capitalize())
