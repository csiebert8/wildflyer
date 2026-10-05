"""Write a solved schedule to an Excel workbook: Grid, List and Summary sheets."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .model import RunInput
from .solver import SolveResult

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONTH_FILLS = (PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="B4C6E7"))
WEEKEND_FILL = PatternFill("solid", fgColor="FFF2CC")
NO_GAMES_FILL = PatternFill("solid", fgColor="EDEDED")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DAY_ABBR = ("M", "Tu", "W", "Th", "F", "Sa", "Su")
DAY_NAMES = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")

GRID_FIRST_ROW = 4  # rows 1-3 are month / weekday / day-of-month headers
GRID_FIRST_COL = 2  # column A holds team codes


def write_output(path: str | Path, run: RunInput, result: SolveResult, input_path: str | Path | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if result.has_schedule:
        _grid(wb.active, run, result)
        _list(wb.create_sheet("List"), run, result)
        _summary(wb.create_sheet("Summary"), run, result, input_path)
    else:
        _summary(wb.active, run, result, input_path)
    wb.save(path)
    return path


def text_color(fill_hex: str) -> str:
    """Black or white text, whichever reads better on the fill colour."""
    r, g, b = (int(fill_hex[i:i + 2], 16) for i in (0, 2, 4))
    return "000000" if 0.299 * r + 0.587 * g + 0.114 * b > 140 else "FFFFFF"


# --- Grid ---------------------------------------------------------------------------


def _grid(ws: Worksheet, run: RunInput, result: SolveResult) -> None:
    ws.title = "Grid"
    dates = run.season_dates
    col_of = {d: GRID_FIRST_COL + i for i, d in enumerate(dates)}
    last_date_col = GRID_FIRST_COL + len(dates) - 1
    teams = list(run.teams)
    row_of = {t: GRID_FIRST_ROW + i for i, t in enumerate(teams)}
    games_per_day = Counter(g.date for g in result.games)

    # Header rows: month bands, weekday, day of month.
    ws.cell(1, 1, run.settings.run_name or "Schedule").font = Font(bold=True)
    month_start = 0
    for i, d in enumerate(dates + [None]):
        if d is None or (i > 0 and d.month != dates[i - 1].month):
            first = dates[month_start]
            c1, c2 = GRID_FIRST_COL + month_start, GRID_FIRST_COL + i - 1
            cell = ws.cell(1, c1, first.strftime("%B %Y").upper())
            fill = MONTH_FILLS[first.month % 2]
            for c in range(c1, c2 + 1):
                ws.cell(1, c).fill = fill
            if c2 > c1:
                ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
            cell.font, cell.alignment = Font(bold=True), CENTER
            month_start = i
    ws.cell(2, 1, "Team").font = Font(bold=True)
    for d, c in col_of.items():
        weekday = ws.cell(2, c, DAY_ABBR[d.weekday()])
        day = ws.cell(3, c, d.day)
        for cell in (weekday, day):
            cell.alignment, cell.border = CENTER, BOX
            cell.font = Font(bold=True)
            if not games_per_day[d]:
                cell.fill = NO_GAMES_FILL
            elif d.weekday() >= 4:
                cell.fill = WEEKEND_FILL
        ws.column_dimensions[get_column_letter(c)].width = 5.5

    # Team rows.
    for t, r in row_of.items():
        team = run.teams[t]
        label = ws.cell(r, 1, t)
        label.font = Font(bold=True, color=text_color(team.color))
        label.fill = PatternFill("solid", fgColor=team.color)
        label.alignment, label.border = CENTER, BOX
        for c in range(GRID_FIRST_COL, last_date_col + 1):
            ws.cell(r, c).border = BOX
    for g in result.games:
        c = col_of[g.date]
        home = ws.cell(row_of[g.home], c, g.away)
        color = run.teams[g.home].color
        home.fill = PatternFill("solid", fgColor=color)
        home.font = Font(bold=True, color=text_color(color), size=9)
        away = ws.cell(row_of[g.away], c, f"@{g.home}")
        away.font = Font(size=9)
        home.alignment = away.alignment = CENTER

    # Games-per-day row under the teams.
    totals_row = GRID_FIRST_ROW + len(teams)
    ws.cell(totals_row, 1, "Games").font = Font(italic=True)
    for d, c in col_of.items():
        if games_per_day[d]:
            cell = ws.cell(totals_row, c, games_per_day[d])
            cell.alignment, cell.font = CENTER, Font(italic=True, size=9)

    # Per-team totals to the right of the dates.
    home_counts = Counter(g.home for g in result.games)
    away_counts = Counter(g.away for g in result.games)
    first = last_date_col + 2
    for offset, title in enumerate(("Games", "Home", "Away")):
        cell = ws.cell(3, first + offset, title)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER
        ws.column_dimensions[get_column_letter(first + offset)].width = 7
    for t, r in row_of.items():
        for offset, value in enumerate((home_counts[t] + away_counts[t], home_counts[t], away_counts[t])):
            ws.cell(r, first + offset, value).alignment = CENTER

    legend = totals_row + 2
    ws.cell(legend, 1, "Key:").font = Font(bold=True)
    ws.cell(legend, 2, "UTA").fill = PatternFill("solid", fgColor="D9D9D9")
    ws.cell(legend, 3, "= home game vs UTA (cell in the home team's colour)")
    ws.cell(legend + 1, 2, "@UTA")
    ws.cell(legend + 1, 3, "= away game at UTA")
    ws.cell(legend + 2, 2).fill = WEEKEND_FILL
    ws.cell(legend + 2, 3, "= Friday-Sunday")
    ws.cell(legend + 3, 2).fill = NO_GAMES_FILL
    ws.cell(legend + 3, 3, "= no games scheduled that day")

    ws.column_dimensions["A"].width = 8
    ws.freeze_panes = ws.cell(GRID_FIRST_ROW, GRID_FIRST_COL)


# --- List -----------------------------------------------------------------------------


def _list(ws: Worksheet, run: RunInput, result: SolveResult) -> None:
    columns = ("Date", "Day", "Away", "Home", "Matchup")
    ws.append(columns)
    for i, width in enumerate((12, 6, 24, 24, 14), start=1):
        ws.cell(1, i).fill, ws.cell(1, i).font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = width
    for g in result.games:
        ws.append([g.date, DAY_NAMES[g.date.weekday()], run.teams[g.away].name, run.teams[g.home].name,
                   f"{g.away} @ {g.home}"])
        ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"


# --- Summary ----------------------------------------------------------------------------


def _summary(ws: Worksheet, run: RunInput, result: SolveResult, input_path: str | Path | None) -> None:
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    status_text = {
        "optimal": "Optimal (best possible under the rules)",
        "feasible": "Feasible (valid; time limit reached before proving it is the best)",
        "infeasible": "No schedule possible",
        "unknown": "No schedule found",
        "invalid": "Error",
    }[result.status]
    s = run.settings
    rows = [
        ("Run name", s.run_name),
        ("Input", str(input_path) if input_path else ""),
        ("Created", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Status", status_text),
        ("Season", f"{s.season_start} to {s.season_end}"),
        ("Games scheduled", len(result.games)),
        ("Soft cost", round(result.soft_cost, 2)),
        ("Solve time (s)", round(result.wall_time, 1)),
    ]
    if result.message:
        rows.append(("Note", result.message))
    for label, value in rows:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)

    def table(title: str, header: tuple[str, ...], body: list[tuple]) -> None:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(list(header))
        for i in range(1, len(header) + 1):
            ws.cell(ws.max_row, i).fill, ws.cell(ws.max_row, i).font = HEADER_FILL, HEADER_FONT
        for row in body:
            ws.append(list(row))
        if not body:
            ws.append(["(none)"])

    if result.has_schedule:
        grouped: dict[str, list] = defaultdict(list)
        for p in result.penalties:
            grouped[p.rule.id].append(p)
        body = []
        for rule_id, items in grouped.items():
            rule = items[0].rule
            for p in items:
                body.append((rule_id, rule.type, p.label, p.amount, round(p.cost, 2)))
        table("Soft rule violations", ("Rule", "Type", "Where", "Units", "Cost"), body)

    table("Rules not applied (rule type not supported yet)",
          ("Rule", "Type", "Note"),
          [(r.id, r.type, r.note) for r in result.skipped_rules])
