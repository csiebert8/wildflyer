from collections import Counter
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from wildflyer.cli import main
from wildflyer.loader import load
from wildflyer.model import Lock, Rule, RunInput, Settings, Team
from wildflyer.output import text_color, write_output
from wildflyer.solver import solve

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "ausl_2027.xlsx"
START = date(2027, 6, 1)


def day(n: int) -> date:
    return START + timedelta(n)


def make_run(rules=(), n_teams=4, n_days=14, locks=()) -> RunInput:
    codes = [f"T{i}" for i in range(1, n_teams + 1)]
    teams = {c: Team(c, f"Team {c}", "4472C4") for c in codes}
    settings = Settings(START, day(n_days - 1), "test", time_limit_seconds=20)
    return RunInput(settings, teams, list(rules), list(locks))


def R(type, id="R1", hard=True, weight=0.0, teams=None, run_teams=4, **kw) -> Rule:
    all_teams = frozenset(f"T{i}" for i in range(1, run_teams + 1))
    if teams is None:
        teams = all_teams if type not in ("LEAGUE_BLACKOUT", "GAMES_PER_DAY", "FIXED_GAME") else frozenset()
    kw = {k: (frozenset(v) if isinstance(v, (set, list, tuple)) else v) for k, v in kw.items()}
    return Rule(id=id, type=type, row=2, hard=hard, weight=weight, teams=frozenset(teams), **kw)


def solved(rules=(), **kw):
    result = solve(make_run(rules, **kw), time_limit=20, workers=4)
    assert result.has_schedule, result.status
    return result


def games_of(result, team):
    return [g for g in result.games if team in (g.home, g.away)]


# Every test league plays each pair once at each team's home: 4 teams -> 12 games.
DOUBLE_RR = R("MATCHUP_GAMES", id="RR", role="home", min=1, max=1)


def test_double_round_robin_and_one_game_per_day():
    result = solved([DOUBLE_RR])
    assert len(result.games) == 12
    assert Counter((g.home, g.away) for g in result.games) == {(h, a): 1 for h in ("T1", "T2", "T3", "T4")
                                                                for a in ("T1", "T2", "T3", "T4") if h != a}
    per_team_day = Counter((t, g.date) for g in result.games for t in (g.home, g.away))
    assert max(per_team_day.values()) == 1


def test_team_games_min_max():
    result = solved([R("TEAM_GAMES", min=5, max=5), R("TEAM_GAMES", id="R2", role="home", min=2, max=3)])
    for t in ("T1", "T2", "T3", "T4"):
        assert len(games_of(result, t)) == 5
        assert 2 <= sum(g.home == t for g in result.games) <= 3


def test_matchup_any_role_counts_both_directions():
    result = solved([R("MATCHUP_GAMES", role="any", min=3, max=3)])
    pairs = Counter(frozenset((g.home, g.away)) for g in result.games)
    assert set(pairs.values()) == {3} and len(pairs) == 6


def test_league_blackout():
    blackout = {day(0), day(1), day(5)}
    result = solved([DOUBLE_RR, R("LEAGUE_BLACKOUT", dates=blackout)])
    assert not {g.date for g in result.games} & blackout


def test_league_blackout_days_of_week():
    result = solved([DOUBLE_RR, R("LEAGUE_BLACKOUT", days={0, 1})])  # no Mon/Tue games
    assert all(g.date.weekday() not in (0, 1) for g in result.games)


def test_games_per_day():
    result = solved([DOUBLE_RR, R("GAMES_PER_DAY", max=1)])
    assert max(Counter(g.date for g in result.games).values()) == 1
    result = solved([DOUBLE_RR, R("GAMES_PER_DAY", min=2, max=2, dates={day(3)})])
    assert sum(g.date == day(3) for g in result.games) == 2


class TestAvailability:
    def test_must_home(self):
        result = solved([DOUBLE_RR, R("AVAILABILITY", teams={"T1"}, role="home", option="must",
                                      dates={day(2), day(3)})])
        assert {g.date for g in result.games if g.home == "T1"} >= {day(2), day(3)}

    def test_cannot_play(self):
        dates = {day(i) for i in range(5)}
        result = solved([DOUBLE_RR, R("AVAILABILITY", teams={"T2"}, option="cannot", dates=dates)])
        assert not any(g.date in dates for g in games_of(result, "T2"))

    def test_only_home(self):
        dates = {day(i) for i in range(4, 10)}
        result = solved([DOUBLE_RR, R("AVAILABILITY", teams={"T3"}, role="home", option="only", dates=dates)])
        assert all(g.date in dates for g in result.games if g.home == "T3")

    def test_opponent_filter(self):
        dates = {day(i) for i in range(7)}
        result = solved([DOUBLE_RR, R("AVAILABILITY", teams={"T1"}, opponents={"T2"}, option="cannot",
                                      dates=dates)])
        assert not any(g.date in dates and {g.home, g.away} == {"T1", "T2"} for g in result.games)

    def test_days_filter(self):
        result = solved([DOUBLE_RR, R("AVAILABILITY", teams={"T1"}, role="home", option="only", days={5, 6},
                                      dates=frozenset())])
        assert all(g.date.weekday() in (5, 6) for g in result.games if g.home == "T1")


def test_season_window_first_and_last():
    result = solved([DOUBLE_RR, R("SEASON_WINDOW", option="first", dates={day(2), day(3), day(4)}),
                     R("SEASON_WINDOW", id="R2", option="last", dates={day(9), day(10)})])
    for t in ("T1", "T2", "T3", "T4"):
        dates = sorted(g.date for g in games_of(result, t))
        assert day(2) <= dates[0] <= day(4)
        assert day(9) <= dates[-1] <= day(10)


def test_fixed_game_and_lock():
    result = solved([DOUBLE_RR, R("FIXED_GAME", teams={"T1"}, opponents={"T2"}, dates={day(6)})],
                    locks=[Lock(2, day(8), "T3", "T4")])
    assert any(g == g.__class__(day(6), "T1", "T2") for g in result.games)
    assert any(g.date == day(8) and (g.home, g.away) == ("T3", "T4") for g in result.games)


def test_infeasible():
    result = solve(make_run([DOUBLE_RR, R("LEAGUE_BLACKOUT", dates={day(i) for i in range(14)})]), time_limit=10)
    assert result.status == "infeasible" and not result.games and result.message


class TestSoft:
    def test_soft_rule_broken_only_when_needed(self):
        # Hard: T1 must host on day 0. Soft: nobody plays on day 0 (weight 7).
        rules = [DOUBLE_RR, R("AVAILABILITY", teams={"T1"}, role="home", option="must", dates={day(0)}),
                 R("LEAGUE_BLACKOUT", id="S1", hard=False, weight=7, dates={day(0), day(1)})]
        result = solved(rules)
        assert sum(g.date == day(1) for g in result.games) == 0
        assert [p.rule.id for p in result.penalties] == ["S1"]
        assert result.soft_cost == 7 and result.penalties[0].amount == 1

    def test_weights_pick_cheaper_violation(self):
        # Only one game fits on day 0, but T1 and T2 both want to host then: keep the dearer rule.
        rules = [DOUBLE_RR, R("GAMES_PER_DAY", max=1),
                 R("AVAILABILITY", id="CHEAP", hard=False, weight=1, teams={"T1"}, role="home", option="must",
                   dates={day(0)}),
                 R("AVAILABILITY", id="DEAR", hard=False, weight=10, teams={"T2"}, role="home", option="must",
                   dates={day(0)})]
        result = solved(rules)
        assert any(g.date == day(0) and g.home == "T2" for g in result.games)
        assert [p.rule.id for p in result.penalties] == ["CHEAP"]

    def test_soft_min(self):
        # Soft: 10 home games for T1 is impossible (3 opponents, 1 each) -> 7 units short.
        result = solved([DOUBLE_RR, R("TEAM_GAMES", hard=False, weight=2, teams={"T1"}, role="home", min=10)])
        assert result.penalties[0].amount == 7 and result.soft_cost == 14

    def test_soft_season_window(self):
        rules = [DOUBLE_RR, R("AVAILABILITY", id="H", teams={"T1"}, option="must", dates={day(0)}),
                 R("SEASON_WINDOW", hard=False, weight=3, option="first", dates={day(5), day(6)})]
        result = solved(rules)
        # T1 must play on day 0, as must its opponent: exactly two teams miss the window.
        assert result.soft_cost == 6


def test_unsupported_rules_are_reported_not_applied():
    result = solved([DOUBLE_RR, R("REMATCH_GAP", id="X", n=5)])
    assert [r.id for r in result.skipped_rules] == ["X"]


def test_example_workbook_solves(tmp_path):
    run = load(EXAMPLE).run
    result = solve(run, time_limit=60)
    assert result.has_schedule
    assert len(result.games) == 90
    home = Counter(g.home for g in result.games)
    assert set(home.values()) == {15}
    tex_home = {g.date for g in result.games if g.home == "TEX"}
    assert tex_home == {date(2027, m, d) for m, ds in ((6, (25, 26, 27)), (7, (1, 2, 3, 15, 16, 17, 29, 30, 31)),
                                                      (8, (12, 13, 14))) for d in ds}
    assert not any(g.home == "UTA" and g.date.month == 8 for g in result.games)
    out = write_output(tmp_path / "out.xlsx", run, result)
    assert openpyxl.load_workbook(out).sheetnames == ["Grid", "List", "Summary"]


class TestOutput:
    def test_grid_cells(self, tmp_path):
        run = make_run([DOUBLE_RR])
        result = solve(run, time_limit=20)
        wb = openpyxl.load_workbook(write_output(tmp_path / "o.xlsx", run, result))
        ws = wb["Grid"]
        g = result.games[0]
        col = 2 + (g.date - START).days
        rows = {ws.cell(r, 1).value: r for r in range(4, 8)}
        home_cell = ws.cell(rows[g.home], col)
        assert home_cell.value == g.away
        assert home_cell.fill.fgColor.rgb.endswith("4472C4")
        assert ws.cell(rows[g.away], col).value == f"@{g.home}"
        assert ws.cell(3, col).value == g.date.day
        assert ws.freeze_panes == "B4"
        assert wb["List"].max_row == 13  # header + 12 games

    def test_infeasible_writes_summary_only(self, tmp_path):
        run = make_run([DOUBLE_RR, R("LEAGUE_BLACKOUT", dates={day(i) for i in range(14)})])
        result = solve(run, time_limit=10)
        wb = openpyxl.load_workbook(write_output(tmp_path / "o.xlsx", run, result))
        assert wb.sheetnames == ["Summary"]
        assert "No schedule possible" in [c.value for c in wb["Summary"]["B"]]

    def test_text_color(self):
        assert text_color("FFFFFF") == "000000" and text_color("0000FF") == "FFFFFF"


def test_cli_solve(tmp_path, capsys):
    out = tmp_path / "sched.xlsx"
    assert main(["solve", str(EXAMPLE), "-o", str(out), "--time-limit", "30"]) == 0
    assert out.exists()
    assert "Status: optimal" in capsys.readouterr().out
