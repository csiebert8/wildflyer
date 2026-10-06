"""Preferences, conflict diagnosis, reports (M4) and base runs / comparison (M5)."""

from collections import Counter

import openpyxl
import pytest

from test_solver import DOUBLE_RR, R, START, day, make_run
from wildflyer.cli import main
from wildflyer.loader import load, read_schedule
from wildflyer.model import Game, Location, Lock
from wildflyer.output import write_output
from wildflyer.solver import solve
from wildflyer.template import write_template

WEEKEND = {5, 6}


def solved(rules, **kw):
    result = solve(make_run(rules, **kw), time_limit=20, workers=8)
    assert result.has_schedule, result.status
    return result


class TestDatePreference:
    def test_prefer_weekends_when_possible(self):
        rules = [DOUBLE_RR, R("DATE_PREFERENCE", hard=False, weight=3, option="prefer", days=WEEKEND)]
        result = solved(rules, n_days=28)  # 8 weekend days x 2 games >= 12 games
        assert all(g.date.weekday() in WEEKEND for g in result.games)
        assert result.soft_cost == 0

    def test_prefer_counts_games_outside(self):
        rules = [DOUBLE_RR, R("DATE_PREFERENCE", hard=False, weight=3, option="prefer", days=WEEKEND)]
        result = solved(rules, n_days=14)  # 4 weekend days x 2 games = 8: 4 games must be weekdays
        assert sum(g.date.weekday() not in WEEKEND for g in result.games) == 4
        assert result.soft_cost == 12

    def test_avoid(self):
        avoided = {day(i) for i in range(7)}
        rules = [DOUBLE_RR, R("DATE_PREFERENCE", hard=False, weight=1, option="avoid", dates=avoided)]
        result = solved(rules, n_days=21)
        assert not any(g.date in avoided for g in result.games)

    def test_teams_and_role_filter(self):
        # Only T1's home games are discouraged on weekdays.
        rules = [DOUBLE_RR, R("DATE_PREFERENCE", hard=False, weight=5, option="prefer", days=WEEKEND,
                              teams={"T1"}, role="home")]
        result = solved(rules, n_days=28)
        assert all(g.date.weekday() in WEEKEND for g in result.games if g.home == "T1")


class TestBalance:
    def test_hard(self):
        rules = [DOUBLE_RR, R("BALANCE", role="home", days=WEEKEND, max=0)]
        result = solved(rules, n_days=28)
        counts = Counter(g.home for g in result.games if g.date.weekday() in WEEKEND)
        assert len({counts[t] for t in ("T1", "T2", "T3", "T4")}) == 1

    def test_soft_counts_excess_spread(self):
        # T1 hosts all 3 of its home games on weekend days; everyone else none -> spread 3, Max 1 -> 2 units.
        locks = [Lock(2, day(4), "T1", "T2"), Lock(3, day(5), "T1", "T3"), Lock(4, day(11), "T1", "T4")]
        rules = [DOUBLE_RR, R("AVAILABILITY", teams={"T2", "T3", "T4"}, role="home", option="cannot",
                              days=WEEKEND, dates=frozenset()),
                 R("BALANCE", id="B", hard=False, weight=2, role="home", days=WEEKEND, max=1)]
        result = solved(rules, locks=locks)
        assert [(p.rule.id, p.amount) for p in result.penalties] == [("B", 2)]


class TestConflicts:
    def test_names_the_conflicting_rules(self):
        rules = [DOUBLE_RR,
                 R("AVAILABILITY", id="MUST", teams={"T1"}, role="home", option="must", dates={day(3)}),
                 R("AVAILABILITY", id="CANT", teams={"T1"}, option="cannot", dates={day(2), day(3)}),
                 R("MAX_CONSECUTIVE_GAME_DAYS", id="OTHER", max=3)]
        result = solve(make_run(rules), time_limit=20)
        assert result.status == "infeasible"
        assert sorted(c.split()[0] for c in result.conflicts) == ["CANT", "MUST"]

    def test_lock_in_conflict(self):
        rules = [DOUBLE_RR, R("LEAGUE_BLACKOUT", id="OFF", dates={day(5)})]
        result = solve(make_run(rules, locks=[Lock(7, day(5), "T1", "T2")]), time_limit=20)
        assert any(c.startswith("OFF") for c in result.conflicts)
        assert any("Locks row 7" in c for c in result.conflicts)

    def test_conflicts_in_summary(self, tmp_path):
        run = make_run([DOUBLE_RR, R("LEAGUE_BLACKOUT", id="OFF", dates={day(i) for i in range(14)})])
        result = solve(run, time_limit=20)
        wb = openpyxl.load_workbook(write_output(tmp_path / "o.xlsx", run, result))
        values = [c.value for row in wb["Summary"].iter_rows() for c in row]
        assert "Conflicting hard rules" in values and any(str(v).startswith("OFF") for v in values)


def test_move_labels_name_both_locations():
    locks = [Lock(2, day(0), "T1", "T2"), Lock(3, day(1), "T3", "T2")]
    rule = R("TRAVEL_REST", hard=False, weight=5, teams={"T2"}, from_loc=Location(keywords=frozenset({"any"})),
             to_loc=Location(keywords=frozenset({"any"})), n=1)
    result = solved([DOUBLE_RR, rule], locks=locks)
    assert result.penalties[0].label.startswith(f"T2: T1 → T3 on {day(1)}")


class TestReports:
    def test_rules_and_checks_sheets(self, tmp_path):
        rules = [DOUBLE_RR, R("GAMES_PER_DAY", id="SOFT", hard=False, weight=1, max=0, dates={day(0)}),
                 R("AVAILABILITY", id="FORCE", teams={"T1"}, option="must", dates={day(0)})]
        run = make_run(rules)
        result = solve(run, time_limit=20)
        wb = openpyxl.load_workbook(write_output(tmp_path / "o.xlsx", run, result))
        rows = {r[0]: r for r in wb["Rules"].iter_rows(min_row=2, values_only=True) if r[0]}
        assert rows["RR"][4] == "Met" and rows["FORCE"][4] == "Met"
        assert rows["SOFT"][4] == "Broken" and rows["SOFT"][5] == 1
        checks = [r for r in wb["Checks"].iter_rows(values_only=True)]
        team_rows = [r for r in checks if r[0] in ("T1", "T2", "T3", "T4") and isinstance(r[1], int)]
        assert {r[1] for r in team_rows[:4]} == {6}  # 6 games each


# --- M5: base runs -------------------------------------------------------------------------

TEAMS = [("T1", "One", ""), ("T2", "Two", ""), ("T3", "Three", ""), ("T4", "Four", "")]
RR_ROW = {"ID": "RR", "Type": "MATCHUP_GAMES", "Role": "home", "Min": 1, "Max": 1}


def workbook(tmp_path, name, settings=None, rules=(RR_ROW,)):
    base = {"Season start": START, "Season end": day(13), "Time limit (seconds)": 10}
    return write_template(tmp_path / name, settings={**base, **(settings or {})}, teams=TEAMS, rules=list(rules))


@pytest.fixture
def base_schedule(tmp_path):
    src = workbook(tmp_path, "base_rules.xlsx")
    out = tmp_path / "base_schedule.xlsx"
    assert main(["solve", str(src), "-o", str(out)]) == 0
    return out


def test_read_schedule_roundtrip(base_schedule):
    games = read_schedule(base_schedule)
    assert len(games) == 12 and all(isinstance(g, Game) for g in games)


def test_lock_base_before(tmp_path, base_schedule):
    base = read_schedule(base_schedule)
    cutoff = day(7)
    src = workbook(tmp_path, "next.xlsx", {"Base run": base_schedule.name, "Lock base before": cutoff})
    result = load(src)
    assert result.ok, [str(i) for i in result.issues]
    run = result.run
    assert {(l.date, l.home, l.away) for l in run.locks} == {(g.date, g.home, g.away) for g in base if g.date < cutoff}
    solved_ = solve(run, time_limit=10)
    early = {g for g in solved_.games if g.date < cutoff}
    assert early == {g for g in base if g.date < cutoff}


def test_change_weight_keeps_base_games(tmp_path, base_schedule):
    # A new rule bans day 0; with a change weight only the games it forces off day 0 should move.
    base = read_schedule(base_schedule)
    blackout = {"ID": "OFF", "Type": "LEAGUE_BLACKOUT", "Dates": START.strftime("%m/%d")}
    src = workbook(tmp_path, "next.xlsx", {"Base run": base_schedule.name, "Change weight": 1},
                   rules=(RR_ROW, blackout))
    run = load(src).run
    result = solve(run, time_limit=10)
    moved_by_rule = sum(g.date == START for g in base)
    kept = len(set(base) & set(result.games))
    assert kept == len(base) - moved_by_rule
    out = write_output(tmp_path / "next_schedule.xlsx", run, result)
    wb = openpyxl.load_workbook(out)
    assert "Changes" in wb.sheetnames
    assert any(r[0] == "BASE" for r in wb["Rules"].iter_rows(values_only=True))


def test_base_run_errors(tmp_path):
    src = workbook(tmp_path, "x.xlsx", {"Base run": "missing.xlsx"})
    assert any("not found" in i.message for i in load(src).errors)
    src = workbook(tmp_path, "y.xlsx", {"Lock base before": day(3)})
    assert any("needs a Base run" in i.message for i in load(src).errors)


def test_compare_cli(tmp_path, base_schedule, capsys):
    src = workbook(tmp_path, "other.xlsx", {"Run name": "other"},
                   rules=(RR_ROW, {"ID": "OFF", "Type": "LEAGUE_BLACKOUT", "Dates": START.strftime("%m/%d")}))
    other = tmp_path / "other_schedule.xlsx"
    assert main(["solve", str(src), "-o", str(other)]) == 0
    capsys.readouterr()
    assert main(["compare", str(base_schedule), str(other)]) == 0
    out = capsys.readouterr().out
    removed, added = (len(set(read_schedule(base_schedule)) - set(read_schedule(other))),
                      len(set(read_schedule(other)) - set(read_schedule(base_schedule))))
    game_lines = [line for line in out.splitlines() if line.startswith(("  + ", "  - ")) and "@" in line]
    assert f"{removed} game(s) only in" in out
    assert sum(line.startswith("  + ") for line in game_lines) == added


def test_solve_without_output_makes_run_folder(tmp_path, monkeypatch):
    src = workbook(tmp_path, "rules.xlsx", {"Run name": "my run"})
    monkeypatch.chdir(tmp_path)
    assert main(["solve", str(src)]) == 0
    folders = list((tmp_path / "runs").iterdir())
    assert len(folders) == 1 and folders[0].name.endswith("_my_run")
    assert {p.name for p in folders[0].iterdir()} == {"rules.xlsx", "schedule.xlsx"}
